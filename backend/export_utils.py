"""
Utilitários de exportação de relatórios do SAPEE.

Gera arquivos Excel (.xlsx) e PDF a partir de um título e uma tabela
(cabeçalhos + linhas), retornando o conteúdo em bytes para resposta HTTP.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _celula(valor) -> str:
    return "" if valor is None else str(valor)


def gerar_xlsx(titulo: str, headers: list, rows: list, subtitulo: str | None = None) -> bytes:
    """Gera um arquivo Excel (.xlsx) com a tabela informada."""
    wb = Workbook()
    ws = wb.active
    ws.title = (titulo or "Relatorio")[:31]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    linha = 1
    ws.cell(row=linha, column=1, value=titulo).font = Font(bold=True, size=14)
    linha += 1
    if subtitulo:
        ws.cell(row=linha, column=1, value=subtitulo).font = Font(italic=True, size=10)
        linha += 1
    linha += 1

    header_row = linha
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=linha, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    linha += 1

    for row in rows:
        for col, valor in enumerate(row, start=1):
            ws.cell(row=linha, column=col, value=valor)
        linha += 1

    # Largura automática das colunas
    for col in range(1, len(headers) + 1):
        maior = len(_celula(headers[col - 1]))
        for row in rows:
            if col - 1 < len(row):
                maior = max(maior, len(_celula(row[col - 1])))
        ws.column_dimensions[get_column_letter(col)].width = min(maior + 3, 50)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gerar_pdf(titulo: str, headers: list, rows: list, subtitulo: str | None = None) -> bytes:
    """Gera um arquivo PDF (paisagem) com a tabela informada."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=titulo,
    )

    styles = getSampleStyleSheet()
    story = [Paragraph(titulo, styles["Title"])]
    if subtitulo:
        story.append(Paragraph(subtitulo, styles["Normal"]))
    story.append(Spacer(1, 0.4 * cm))

    dados = [headers] + [[_celula(v) for v in row] for row in rows]
    tabela = Table(dados, repeatRows=1)
    tabela.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(tabela)
    doc.build(story)
    return buf.getvalue()
