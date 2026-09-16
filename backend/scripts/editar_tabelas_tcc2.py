# --- Bootstrap de execucao (injetado): garante imports do backend e carrega o .env ---
import os as _os
import sys as _sys

_BACKEND_DIR = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
if _BACKEND_DIR not in _sys.path:
    _sys.path.insert(0, _BACKEND_DIR)
try:
    from dotenv import load_dotenv as _load_dotenv

    _load_dotenv(_os.path.join(_BACKEND_DIR, ".env"))
except Exception:
    pass
# --- fim do bootstrap ---

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Atualiza a Lista de Tabelas e insere/atualiza tabelas no TCC2.
Salva como TCC2_SAPEE_COMPLETO_V4.docx
"""
import sys
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

INPUT = '../docs/TCC2/TCC2_SAPEE_COMPLETO_V3.docx'
OUTPUT = '../docs/TCC2/TCC2_SAPEE_COMPLETO_V4.docx'

# ---------------------------------------------------------------------------
# Dados das 13 tabelas da Lista de Tabelas
# ---------------------------------------------------------------------------
LISTA_TABELAS = [
    ("Tabela 1", "Taxa de Evasão por Tipo de Instituição (Brasil, 2019–2024)", "[p.]"),
    ("Tabela 2", "Taxa de Evasão no Ensino Superior por Região (2023)", "[p.]"),
    ("Tabela 3", "Pesos dos Fatores de Risco Acadêmicos", "[p.]"),
    ("Tabela 4", "Pesos dos Fatores de Risco Socioeconômicos", "[p.]"),
    ("Tabela 5", "Comparativo entre Algoritmos de Machine Learning Testados", "[p.]"),
    ("Tabela 6", "Matriz de Confusão do Modelo Final (Random Forest)", "[p.]"),
    ("Tabela 7", "Métricas de Desempenho do Modelo Preditivo", "[p.]"),
    ("Tabela 8", "Distribuição dos 2.700 Estudantes por Nível de Risco", "[p.]"),
    ("Tabela 9", "Dimensões e Variáveis do Questionário Psicossocial", "[p.]"),
    ("Tabela 10", "Categorias de Intervenções Pedagógicas por Nível de Risco", "[p.]"),
    ("Tabela 11", "Comparação com Outros Sistemas de Alerta Precoce", "[p.]"),
    ("Tabela 12", "Tecnologias Utilizadas no Desenvolvimento do SAPEE", "[p.]"),
    ("Tabela 13", "Cronograma de Execução do Projeto", "[p.]"),
]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_cell_text(cell, text, bold=False, font_size=10):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.bold = bold
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def fill_table(table, rows, bold_header=True, font_size=9):
    """Preenche uma tabela do docx com dados (lista de listas)."""
    # Limpa linhas existentes
    for row in table.rows:
        for cell in row.cells:
            cell.text = ''
    # Ajusta número de linhas
    while len(table.rows) < len(rows):
        table.add_row()
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        while len(row.cells) < len(row_data):
            row._tr.append(deepcopy(row.cells[-1]._tc))
        for j, val in enumerate(row_data):
            set_cell_text(row.cells[j], str(val), bold=(bold_header and i == 0), font_size=font_size)


def find_paragraph_after(doc, start_idx, text_contains):
    """Retorna índice do primeiro parágrafo após start_idx que contém text_contains."""
    for i in range(start_idx, len(doc.paragraphs)):
        if text_contains.lower() in doc.paragraphs[i].text.lower():
            return i
    return None


def insert_paragraph_after(doc, paragraph_idx, text, style='Normal'):
    """Insere um novo parágrafo após paragraph_idx."""
    # Adiciona novo parágrafo no final
    new_para = doc.add_paragraph(text, style=style)
    # Move para após o parágrafo alvo
    target_p = doc.paragraphs[paragraph_idx]._p
    target_p.addnext(new_para._p)
    return new_para


def insert_table_after(doc, paragraph_idx, rows, bold_header=True, font_size=9):
    """Insere tabela após o parágrafo paragraph_idx."""
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = 'Table Grid'
    fill_table(table, rows, bold_header, font_size)
    # Move tabela para após o parágrafo
    target_p = doc.paragraphs[paragraph_idx]._p
    target_p.addnext(table._tbl)
    return table


# ---------------------------------------------------------------------------
# Abrir documento
# ---------------------------------------------------------------------------
doc = Document(INPUT)
print(f'Aberto: {INPUT}')

# ---------------------------------------------------------------------------
# 1. Atualizar Lista de Tabelas (Table 1)
# ---------------------------------------------------------------------------
tabela_lista = doc.tables[1]
fill_table(tabela_lista, [
    ["Tabela", "Descrição", "Página"]
] + [[n, d, p] for n, d, p in LISTA_TABELAS], bold_header=True, font_size=10)
print('Lista de Tabelas atualizada com 13 itens.')

# ---------------------------------------------------------------------------
# 2. Corrigir Tabela 2: Ensino Técnico -> Ensino Superior
# ---------------------------------------------------------------------------
# Localizar o parágrafo do caption
for i, p in enumerate(doc.paragraphs):
    if 'Tabela 2 - Taxa de Evasão no Ensino Técnico' in p.text:
        p.text = 'Tabela 2 - Taxa de Evasão no Ensino Superior por Região (Brasil, 2023)'
        print(f'Caption Tabela 2 corrigido no parágrafo {i}')
        break

# Atualizar dados da Tabela 2 (Table 4)
tabela2 = doc.tables[4]
fill_table(tabela2, [
    ["Região", "Ensino Superior Público (%)", "Ensino Superior Privado (%)", "Média Regional (%)"],
    ["Norte", "19,8", "22,4", "21,1"],
    ["Nordeste", "18,2", "21,7", "19,9"],
    ["Centro-Oeste", "15,6", "18,9", "17,2"],
    ["Sudeste", "14,3", "17,5", "15,9"],
    ["Sul", "13,9", "16,8", "15,3"],
], bold_header=True, font_size=9)
print('Conteúdo da Tabela 2 atualizado para ensino superior.')

# ---------------------------------------------------------------------------
# 3. Atualizar cronograma (Quadro 4.2 -> Tabela 13) com data 10/01/2026
# ---------------------------------------------------------------------------
for i, p in enumerate(doc.paragraphs):
    if 'Quadro 4.2' in p.text and 'Sprint' in p.text:
        p.text = 'Tabela 13 - Cronograma de Execução do Projeto (início: 10 de janeiro de 2026)'
        print(f'Caption cronograma renumerado no parágrafo {i}')
        break

tabela_cronograma = doc.tables[14]
fill_table(tabela_cronograma, [
    ["Sprint", "Duração", "Título", "Entregas Principais", "Período (2026)"],
    ["Sprint 0", "2 semanas", "Planejamento e Setup", "Definição do escopo; elicitação inicial de requisitos; configuração do ambiente Railway/Docker; setup do repositório Git", "10/01 a 23/01"],
    ["Sprint 1", "2 semanas", "Fundação Backend", "Modelos SQLAlchemy; autenticação JWT; CRUD de usuários e cursos", "24/01 a 06/02"],
    ["Sprint 2", "2 semanas", "Gestão de Alunos e Importação", "Cadastro e gestão de alunos; importação CSV; modelagem de dados acadêmicos", "07/02 a 20/02"],
    ["Sprint 3", "2 semanas", "Frontend Base e Dashboard", "React + TypeScript + TailwindCSS; layout responsivo; dashboard inicial", "21/02 a 06/03"],
    ["Sprint 4", "2 semanas", "Motor de Predição", "Pipeline ML; treinamento e validação do Random Forest; geração de predições", "07/03 a 20/03"],
    ["Sprint 5", "2 semanas", "Intervenções e Notificações", "Módulo de intervenções pedagógicas; notificações Telegram; planos de ação", "21/03 a 03/04"],
    ["Sprint 6", "2 semanas", "Questionário e Relatórios", "Questionário psicossocial; tokens; relatórios gerenciais e indicadores", "04/04 a 17/04"],
    ["Sprint 7", "2 semanas", "Validação e Ajustes", "Testes funcionais; validação com egressos; ajustes no modelo; documentação TCC2", "18/04 a 01/05"],
], bold_header=True, font_size=9)
print('Cronograma atualizado com início em 10/01/2026.')

# ---------------------------------------------------------------------------
# 4. Inserir Tabela 12: Tecnologias Utilizadas
# ---------------------------------------------------------------------------
# Procurar seção 3.6 ARQUITETURA DO SISTEMA ou 4.1 REQUISITOS
secao_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper().startswith('3.6 ARQUITETURA DO SISTEMA'):
        secao_idx = i
        break

if secao_idx:
    # Procurar primeiro parágrafo após a seção que mencione tecnologia/stack
    insert_idx = None
    for i in range(secao_idx + 1, min(secao_idx + 40, len(doc.paragraphs))):
        if any(k in doc.paragraphs[i].text.lower() for k in ['tecnologia', 'stack', 'ferramenta', 'fastapi', 'react']):
            insert_idx = i
            break
    if insert_idx is None:
        insert_idx = secao_idx + 2

    insert_paragraph_after(doc, insert_idx, 'Tabela 12 - Tecnologias Utilizadas no Desenvolvimento do SAPEE', style='Caption')
    tabela_tec = insert_table_after(doc, insert_idx + 1, [
        ["Camada", "Tecnologia", "Versão", "Função no SAPEE"],
        ["Backend", "Python + FastAPI", "3.12 / 0.109+", "API REST, autenticação JWT, lógica de negócio e integração ML"],
        ["Frontend", "React + TypeScript + Vite", "19 / 5.8 / 6", "SPA responsiva, dashboards, formulários e visualização de dados"],
        ["Banco de Dados", "MySQL + SQLAlchemy ORM", "8.0 / 2.0", "Persistência relacional dos dados acadêmicos e preditivos"],
        ["Machine Learning", "Scikit-learn + Pandas + NumPy", "1.4", "Pipeline de predição, pré-processamento e métricas"],
        ["Estilo", "TailwindCSS", "4", "Estilização utilitária e responsiva do frontend"],
        ["Gráficos", "Recharts", "3", "Visualizações interativas de risco e indicadores"],
        ["Infraestrutura", "Docker Compose + Railway", "-", "Containerização e deploy na nuvem (sapee-dewas.up.railway.app)"],
        ["Notificações", "python-telegram-bot + SMTP", "-", "Envio de alertas via Telegram e e-mail"],
    ], bold_header=True, font_size=9)
    print(f'Tabela 12 inserida após parágrafo {insert_idx}')
else:
    print('Seção 3.6 não encontrada; Tabela 12 não inserida.')

# ---------------------------------------------------------------------------
# 5. Ajustar captions de tabelas já existentes para numeração da lista
# ---------------------------------------------------------------------------
# Tabela 6 (Matriz de Confusão) = Table 38 (TABELA 8 no cap 5)
# Tabela 7 (Métricas) = Table 39 (TABELA 9 no cap 5)
# Tabela 8 (Distribuição por risco) = Table 35 (TABELA 5 no cap 5)
# Tabela 11 (Comparação com sistemas) = Table 42 (TABELA 12 no cap 5)

for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'TABELA 8: Matriz de Confusão do Modelo Preditivo SAPEE (n = 225)':
        p.text = 'Tabela 6 - Matriz de Confusão do Modelo Final (Random Forest) (n = 225)'
        print(f'Renum Tabela 6 no parágrafo {i}')
    elif p.text.strip() == 'TABELA 9: Métricas de Desempenho do Modelo Preditivo (n = 225)':
        p.text = 'Tabela 7 - Métricas de Desempenho do Modelo Preditivo (n = 225)'
        print(f'Renum Tabela 7 no parágrafo {i}')
    elif p.text.strip() == 'TABELA 5: Distribuição de Alunos por Nível de Risco':
        p.text = 'Tabela 8 - Distribuição dos 2.700 Estudantes por Nível de Risco'
        print(f'Renum Tabela 8 no parágrafo {i}')
    elif p.text.strip() == 'TABELA 12: Comparação do SAPEE com Sistemas Referenciados na Literatura':
        p.text = 'Tabela 11 - Comparação do SAPEE com Outros Sistemas de Alerta Precoce'
        print(f'Renum Tabela 11 no parágrafo {i}')

# ---------------------------------------------------------------------------
# Salvar
# ---------------------------------------------------------------------------
doc.save(OUTPUT)
print(f'\nDocumento salvo em: {OUTPUT}')
