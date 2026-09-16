# --- Bootstrap de execucao (injetado): garante imports do backend e carrega o .env ---
import os as _os
import sys as _sys

_BACKEND_DIR = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
if _BACKEND_DIR not in _sys.path:
    _sys.path.insert(0, _BACKEND_DIR)
try:
    from dotenv import load_dotenv as _load_dotenv

    _load_dotenv(_os.path.join(_BACKEND_DIR, ".env"))
except Exception:
    pass
# --- fim do bootstrap ---

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Atualiza a Lista de Tabelas incluindo TODAS as tabelas/quadros do documento.
Salva como TCC2_SAPEE_COMPLETO_V6.docx
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt

INPUT = '../docs/TCC2/TCC2_SAPEE_COMPLETO_V5.docx'
OUTPUT = '../docs/TCC2/TCC2_SAPEE_COMPLETO_V6.docx'

PARAGRAFOS_POR_PAGINA = 22.0

def estimar_pagina(idx):
    return max(1, int(idx / PARAGRAFOS_POR_PAGINA) + 1)

doc = Document(INPUT)

# ---------------------------------------------------------------------------
# Tabelas já principais (13)
# ---------------------------------------------------------------------------
tabelas_principais = [
    (536, "Taxa de Evasão por Tipo de Instituição (Brasil, 2019–2024)"),
    (555, "Taxa de Evasão no Ensino Superior por Região (2023)"),
    (569, "Pesos dos Fatores de Risco Acadêmicos"),
    (576, "Pesos dos Fatores de Risco Socioeconômicos"),
    (622, "Comparativo entre Algoritmos de Machine Learning Testados"),
    (1823, "Matriz de Confusão do Modelo Final (Random Forest)"),
    (1831, "Métricas de Desempenho do Modelo Preditivo"),
    (1799, "Distribuição dos 2.700 Estudantes por Nível de Risco"),
    (647, "Dimensões e Variáveis do Questionário Psicossocial"),
    (659, "Categorias de Intervenções Pedagógicas por Nível de Risco"),
    (1865, "Comparação do SAPEE com Outros Sistemas de Alerta Precoce"),
    (406, "Tecnologias Utilizadas no Desenvolvimento do SAPEE"),
    (792, "Cronograma de Execução do Projeto"),
]

# ---------------------------------------------------------------------------
# Tabelas complementares do capítulo 3 e 4
# ---------------------------------------------------------------------------
tabelas_complementares = [
    (753, "Tabela 3.1 - Comparativo de Funcionalidades entre Sistemas de Predição de Evasão"),
    (760, "Tabela 3.2 - Análise SWOT do SAPEE"),
    (771, "Quadro 3.1 - Síntese da Revisão de Trabalhos Relacionados"),
    (785, "Quadro 4.1 - Classificação Metodológica da Pesquisa"),
    (810, "Tabela 4.1 - Requisitos Funcionais do SAPEE"),
    (816, "Tabela 4.2 - Requisitos Não Funcionais do SAPEE"),
    (887, "Quadro 4.3 - Especificação de Caso de Uso: Executar Predição de Risco"),
    (948, "Tabela 4.3 - Tabelas do Modelo de Dados do SAPEE"),
    (1050, "Quadro 4.4 - Métricas-Alvo do Modelo Preditivo do SAPEE"),
]

# ---------------------------------------------------------------------------
# Tabelas do capítulo 4/5 sobre API, algoritmo e frontend
# ---------------------------------------------------------------------------
tabelas_api_algo = [
    (1208, "TABELA 1 - Endpoints da API REST do SAPEE - Mapeamento Completo por Módulo"),
    (1236, "TABELA 2 - Papéis de Usuário (Roles) e Permissões do SAPEE"),
    (1322, "TABELA 3 - Fatores de Risco e Seus Pesos Individuais no Algoritmo v2"),
    (1352, "TABELA 4 - Combinações Perigosas e Seus Multiplicadores"),
    (1376, "TABELA 5 - Limiares de Classificação de Risco do SAPEE"),
    (1394, "TABELA 6 - Correspondência entre Fator Detectado e Recomendação Gerada"),
    (1400, "TABELA 7 - Dimensões do Questionário Psicossocial e Pesos no Cálculo"),
    (1455, "TABELA 8 - Páginas do Frontend SAPEE"),
    (1460, "TABELA 9 - Principais Componentes da Arquitetura Frontend"),
    (1573, "TABELA 10 - Canais de Comunicação Suportados pelo SAPEE"),
]

# ---------------------------------------------------------------------------
# Tabelas do capítulo 5 de resultados (excluindo as já listadas nas 13 principais)
# ---------------------------------------------------------------------------
tabelas_resultados = [
    (1778, "TABELA 1 - Distribuição de Alunos por Curso de Graduação"),
    (1782, "TABELA 2 - Composição dos Egressos por Motivo de Saída"),
    (1789, "TABELA 3 - Distribuição de Alunos e Egressos por Coorte"),
    (1794, "TABELA 4 - Indicadores do Dashboard SAPEE - Visão Consolidada"),
    (1806, "TABELA 6 - Distribuição de Risco por Curso"),
    (1813, "TABELA 7 - Status das Intervenções Registradas"),
    (1837, "TABELA 10 - Métricas de Desempenho por Nível de Risco (n = 225)"),
    (1843, "TABELA 11 - Métricas de Desempenho por Coorte (Ano de Ingresso)"),
    (1879, "TABELA 13 - Resumo de Testes Funcionais por Módulo"),
]

# ---------------------------------------------------------------------------
# Montar lista completa
# ---------------------------------------------------------------------------
todas = tabelas_principais + tabelas_complementares + tabelas_api_algo + tabelas_resultados
print(f'Total de tabelas a incluir na Lista de Tabelas: {len(todas)}')

# ---------------------------------------------------------------------------
# Atualizar Lista de Tabelas (doc.tables[1])
# ---------------------------------------------------------------------------
tabela_lista = doc.tables[1]
rows = [["Tabela", "Descrição", "Página"]]
for n, (idx, desc) in enumerate(todas, 1):
    pag = str(estimar_pagina(idx))
    rows.append([f"Tabela {n}", desc, pag])

# Limpa e recria
for row in tabela_lista.rows:
    for cell in row.cells:
        cell.text = ''
while len(tabela_lista.rows) < len(rows):
    tabela_lista.add_row()
for i, row_data in enumerate(rows):
    row = tabela_lista.rows[i]
    for j, val in enumerate(row_data):
        cell = row.cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.size = Pt(10)
        run.bold = (i == 0)
        p.alignment = 1  # CENTER

print('Lista de Tabelas atualizada com todas as tabelas do documento.')

# ---------------------------------------------------------------------------
# Salvar
# ---------------------------------------------------------------------------
doc.save(OUTPUT)
print(f'\nDocumento salvo em: {OUTPUT}')
