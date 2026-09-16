# --- Bootstrap de execucao (injetado): garante imports do backend e carrega o .env ---
import os as _os
import sys as _sys

_BACKEND_DIR = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
if _BACKEND_DIR not in _sys.path:
    _sys.path.insert(0, _BACKEND_DIR)
try:
    from dotenv import load_dotenv as _load_dotenv

    _load_dotenv(_os.path.join(_BACKEND_DIR, ".env"))
except Exception:
    pass
# --- fim do bootstrap ---

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Correções no TCC2 V4: atualiza páginas de figuras/tabelas e referências cruzadas.
Salva como V5.
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt

INPUT = '../docs/TCC2/TCC2_SAPEE_COMPLETO_V4.docx'
OUTPUT = '../docs/TCC2/TCC2_SAPEE_COMPLETO_V5.docx'

PARAGRAFOS_POR_PAGINA = 22.0

def estimar_pagina(idx):
    return max(1, int(idx / PARAGRAFOS_POR_PAGINA) + 1)

doc = Document(INPUT)

# ---------------------------------------------------------------------------
# 1. Mapear captions de figuras para posição da imagem
# ---------------------------------------------------------------------------
fig_pos = {}
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Figura ') and (' - ' in txt or ':' in txt):
        try:
            num = int(txt.split('Figura')[1].split()[0].split('-')[0].split(':')[0].strip())
            # encontrar imagem mais próxima
            img_pos = None
            for j in range(max(0, i-3), min(len(doc.paragraphs), i+2)):
                if bool(doc.paragraphs[j]._element.xpath('.//w:drawing')):
                    img_pos = j
                    break
            fig_pos[num] = img_pos if img_pos is not None else i
        except:
            pass
    elif '[INSERIR PRINTSCREEN: Figura ' in txt:
        try:
            num = int(txt.split('Figura')[1].split()[0].split('-')[0].strip())
            if num not in fig_pos:
                fig_pos[num] = i
        except:
            pass

# ---------------------------------------------------------------------------
# 2. Mapear captions de tabelas principais
# ---------------------------------------------------------------------------
tab_pos = {}
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Tabela ') and (' - ' in txt or ':' in txt):
        try:
            num = int(txt.split('Tabela')[1].split()[0].split('-')[0].split(':')[0].strip())
            tab_pos[num] = i
        except:
            pass

# ---------------------------------------------------------------------------
# 3. Atualizar Lista de Figuras
# ---------------------------------------------------------------------------
print('--- Atualizando Lista de Figuras ---')
for row in doc.tables[0].rows[1:]:
    cells = row.cells
    num_str = cells[0].text.strip()
    desc = cells[1].text.strip()
    try:
        num = int(num_str.replace('Figura ', '').strip())
    except:
        continue
    if num in fig_pos:
        idx = fig_pos[num]
        has_img = bool(doc.paragraphs[idx]._element.xpath('.//w:drawing')) if idx < len(doc.paragraphs) else False
        if not has_img:
            # placeholder, manter [p.]
            new_pag = '[p.]'
        else:
            new_pag = str(estimar_pagina(idx))
        if cells[2].text.strip() != new_pag:
            print(f'Figura {num}: {cells[2].text} -> {new_pag}')
            cells[2].text = new_pag
            for p in cells[2].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

# ---------------------------------------------------------------------------
# 4. Atualizar Lista de Tabelas
# ---------------------------------------------------------------------------
print('--- Atualizando Lista de Tabelas ---')
for row in doc.tables[1].rows[1:]:
    cells = row.cells
    num_str = cells[0].text.strip()
    try:
        num = int(num_str.replace('Tabela ', '').strip())
    except:
        continue
    if num in tab_pos:
        idx = tab_pos[num]
        new_pag = str(estimar_pagina(idx))
        if cells[2].text.strip() != new_pag:
            print(f'Tabela {num}: {cells[2].text} -> {new_pag}')
            cells[2].text = new_pag
            for p in cells[2].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

# ---------------------------------------------------------------------------
# 5. Corrigir referência cruzada: dashboard (Figura 2) -> (Figura 1)
# ---------------------------------------------------------------------------
print('--- Corrigindo referências cruzadas ---')
for i, p in enumerate(doc.paragraphs):
    txt = p.text
    if 'dashboard' in txt.lower() and 'Figura 2' in txt:
        p.text = txt.replace('Figura 2', 'Figura 1')
        print(f'Parágrafo {i}: corrigido "Figura 2" -> "Figura 1" no contexto de dashboard')

# ---------------------------------------------------------------------------
# Salvar
# ---------------------------------------------------------------------------
doc.save(OUTPUT)
print(f'\nDocumento corrigido salvo em: {OUTPUT}')
